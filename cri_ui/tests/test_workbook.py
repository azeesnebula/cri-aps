"""
Header detection for planner-uploaded planning workbooks; no database.

The fixtures below are deliberately small but reproduce the three ways the
real workbooks differ from each other, because those differences are exactly
what a fixed-offset reader gets wrong:

  * a leading title row that is present in one workbook and not the other
  * a legend block of a dozen rows before the header
  * an extra column (QTY) that shifts every later column along by one
"""

import io
import os
import unittest

from openpyxl import Workbook

from cri_ui.workbook import SiegerWorkbook, locateHeader, norm, SHEETS


def spec(key):
    return next(s for s in SHEETS if s.key == key)


def book(sheets):
    """sheets: {name: [row, row, ...]} -> an in-memory .xlsx"""
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(title=name)
        for row in rows:
            ws.append(list(row))
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# The seven sheets a complete workbook needs, in the "no title row" layout.
def minimalSheets(with_qty=False, item_legend=0, machine_title=0):
    wo_human = ["Work order id", "Routing ID", "Unit of measurement", "Item"]
    wo_code = ["NOF", "NOMG", "UTIL", "CODEARTIC"]
    wo_data = ["WO1", "RT1", "NOS", "ITEM1"]
    if with_qty:
        # QTY sits between the routing and the unit of measure, with a blank
        # human header above it - exactly as in APS Sample Data.xlsx
        wo_human = ["Work order id", "Routing ID", "", "Unit of measurement", "Item"]
        wo_code = ["NOF", "NOMG", "QTY", "UTIL", "CODEARTIC"]
        wo_data = ["WO1", "RT1", 7, "NOS", "ITEM1"]

    sheets = {
        "9.Sections within plant": [
            ["Section ID", "Section designation"],
            ["CODESECTI", "DESIGSECT"],
            ["HFAB", "HEAVY FABRICATION"],
        ],
        "5.Machine groupings": [
            ["Machine grouping", "Machine grouping designation", "Operating mode"],
            ["ILOT", "LIBILOT", "MODEILOT"],
            ["SMMG1", "MIG WELDING", ""],
        ],
        "6.Machines list": (
            [["Machines - that is, primary resources"]] * machine_title
            + [
                ["Machine grouping", "Machine", "Section ID", "Calendar ID", "Machine designation"],
                ["ILOT", "MACHINE", "CODESECTI", "NOCALHEBD", "LIBMACH"],
                ["MCDR1", "DRL-B41-05", "MCSP", "", "DRILLING MACHINE - 5"],
            ]
        ),
        "3.list of Operations": [
            ["Operation", "Machine grouping", "Operation Qty", "Operation Cycle time", "Operation designation"],
            ["OPE", "ILOT", "CODEBASET", "DURREAL", "LIBOP"],
            ["FABR", "SMMG1", "1", "2.5", "SMD FABRICATION"],
        ],
        "7.Phases of the routings ": [
            ["Routing ID", "Phase code", "Phase code", "Phase designation", "Min"],
            ["NOMG", "NOPHASE", "OPE", "LIBPHASE", "Cycle Time", "Setup Time"],
            ["RT1", 10, "FABR", "SMD FABRICATION", 2, 5],
        ],
        "2.List of all the items": (
            [["", "", "", "BO - Boughtout"]] * item_legend
            + [
                ["Item", "Unit of measurement", "Designation/Description", "Material type"],
                ["CODEARTIC", "UTIL", "LIBARTIC", "TYPEMATI"],
                ["ITEM1", "NOS", "A PRODUCT", "SF"],
            ]
        ),
        "11.Work order list": [wo_human, wo_code, wo_data],
    }
    return sheets


class NormaliseTest(unittest.TestCase):
    def test_strips_everything_but_alphanumerics(self):
        self.assertEqual(norm(" Machine  grouping "), "machinegrouping")
        self.assertEqual(norm("Designation/Description"), "designationdescription")
        self.assertEqual(norm("11.Work order list"), "11workorderlist")
        self.assertEqual(norm(None), "")


class HeaderDetectionTest(unittest.TestCase):
    def test_finds_columns_and_first_data_row(self):
        rows = [
            ["Machines - that is, primary resources", None, None],
            ["Machine grouping", "Machine", "Section ID"],
            ["ILOT", "MACHINE", "CODESECTI"],
            ["MCDR1", "DRL-01", "MCSP"],
        ]
        cols, first = locateHeader(rows, spec("machines"))
        self.assertEqual(cols["group"], 0)
        self.assertEqual(cols["machine"], 1)
        self.assertEqual(cols["section"], 2)
        self.assertEqual(first, 3)

    def test_technical_code_wins_over_a_duplicated_label(self):
        # Both column 1 and column 2 are labelled "Phase code"; only the LN
        # codes distinguish the sequence number from the operation code.
        rows = [
            ["Routing ID", "Phase code", "Phase code", "Phase designation"],
            ["NOMG", "NOPHASE", "OPE", "LIBPHASE"],
            ["RT1", 10, "FABR", "FABRICATION"],
        ]
        cols, first = locateHeader(rows, spec("phases"))
        self.assertEqual(cols["seq"], 1)
        self.assertEqual(cols["op"], 2)
        self.assertEqual(first, 2)

    def test_header_below_a_long_legend_block(self):
        rows = [["", "", "", "BO - Boughtout"]] * 12
        rows += [
            ["Item", "Unit of measurement", "Designation/Description"],
            ["CODEARTIC", "UTIL", "LIBARTIC"],
            ["ITEM1", "NOS", "A PRODUCT"],
        ]
        cols, first = locateHeader(rows, spec("items"))
        self.assertEqual(cols["code"], 0)
        self.assertEqual(first, 14)


class WorkbookResolutionTest(unittest.TestCase):
    def test_clean_workbook_has_no_blocking_problems(self):
        wbk = SiegerWorkbook(book(minimalSheets())).read()
        report = wbk.report()
        self.assertEqual(report["blocking"], [])
        for sheet in report["sheets"]:
            self.assertEqual(sheet["rows"], 1, sheet["key"])

    def test_extra_title_and_legend_rows_do_not_shift_the_data(self):
        """The difference between MASTER DATA.XLSX and APS Sample Data.xlsx."""
        plain = SiegerWorkbook(book(minimalSheets())).read()
        padded = SiegerWorkbook(
            book(minimalSheets(item_legend=12, machine_title=1))
        ).read()
        self.assertEqual(padded.report()["blocking"], [])
        for key in ("items", "machines"):
            self.assertEqual(
                len(plain.sheets[key].rows),
                len(padded.sheets[key].rows),
                "%s lost or gained a row when the sheet was padded" % key,
            )
        # ... and the padding really did move the header
        self.assertGreater(
            padded.sheets["items"].first_data_row,
            plain.sheets["items"].first_data_row,
        )

    def test_quantity_column_is_used_when_present(self):
        wbk = SiegerWorkbook(book(minimalSheets(with_qty=True))).read()
        wo = wbk.sheets["workorders"]
        self.assertIn("qty", wo.columns)
        row = wo.rows[0]
        # The real risk: reading positionally puts the uom in the item field
        self.assertEqual(wo.text(row, "item"), "ITEM1")
        self.assertEqual(wo.text(row, "uom"), "NOS")
        self.assertEqual(wo.get(row, "qty"), 7)
        self.assertTrue(
            any("read from the workbook" in n for n in wbk.report()["notes"])
        )

    def test_quantity_absent_is_reported_as_generated(self):
        wbk = SiegerWorkbook(book(minimalSheets(with_qty=False))).read()
        self.assertNotIn("qty", wbk.sheets["workorders"].columns)
        self.assertTrue(
            any("will be generated" in n for n in wbk.report()["notes"])
        )

    def test_missing_sheet_is_reported_not_guessed(self):
        sheets = minimalSheets()
        del sheets["6.Machines list"]
        report = SiegerWorkbook(book(sheets)).read().report()
        self.assertTrue(any("machines" in p for p in report["blocking"]))

    def test_missing_required_column_is_reported(self):
        sheets = minimalSheets()
        # strip the item code column from both header rows
        sheets["2.List of all the items"] = [
            r[1:] for r in sheets["2.List of all the items"]
        ]
        report = SiegerWorkbook(book(sheets)).read().report()
        self.assertTrue(
            any("could not find the column for 'code'" in p for p in report["blocking"])
        )

    def test_unknown_sheets_are_ignored_not_fatal(self):
        sheets = minimalSheets()
        sheets["DSQ Details"] = [["whatever"]]
        wbk = SiegerWorkbook(book(sheets)).read()
        self.assertEqual(wbk.report()["blocking"], [])
        self.assertIn("DSQ Details", wbk.unmatched)

    def test_a_non_spreadsheet_is_refused_clearly(self):
        from cri_ui.workbook import WorkbookProblem

        with self.assertRaises(WorkbookProblem):
            SiegerWorkbook(io.BytesIO(b"this is not a workbook")).read()


class RealWorkbookTest(unittest.TestCase):
    """Runs only where the customer workbooks are checked out."""

    files = (
        "APS Sample Data.xlsx",
        "MASTER DATA.XLSX",
        "1.SAMPLE DATA FOR PLANNING.XLSX",
    )

    def test_every_real_workbook_resolves(self):
        seen = 0
        for name in self.files:
            if not os.path.exists(name):
                continue
            seen += 1
            report = SiegerWorkbook(name).read().report()
            self.assertEqual(report["blocking"], [], "%s: %s" % (name, report["blocking"]))
            rows = {s["key"]: s["rows"] for s in report["sheets"]}
            for key, count in rows.items():
                self.assertGreater(count, 0, "%s: %s is empty" % (name, key))
        if not seen:
            self.skipTest("no customer workbooks in the working directory")


if __name__ == "__main__":
    unittest.main()
