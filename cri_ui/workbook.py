#
# Read a Sieger / Infor-LN planning workbook and turn it into frePPLe rows.
#
# Sheets are located by name and columns by their header text, never by fixed
# row or column offsets. That is not a nicety. The two workbooks we have differ:
#
#   - "MASTER DATA.XLSX" carries one extra leading row in the operations and
#     machines sheets, so an offset tuned to one workbook drops the first data
#     row of the other.
#   - "APS Sample Data.xlsx" has a QTY column in the work order sheet that
#     MASTER DATA.XLSX does not have, shifting every later column by one. Read
#     positionally, every demand ends up with the unit of measure as its item.
#
# Each column is looked up by its Infor-LN technical code first (CODEARTIC,
# NOF, QTY, ...) and only then by the human label above it. The LN codes are
# stable across languages and are unique, while the human labels are not: the
# phases sheet has two different columns both labelled "Phase code".
#

from collections import defaultdict
import re


def norm(value):
    """
    Normalise a header or sheet name for comparison: lowercase, with
    everything that is not a letter or digit removed.
    """
    if value is None:
        return ""
    return re.sub(r"[^a-z0-9]", "", str(value).lower())


class Field:
    """
    One column we need. `code` is the Infor-LN technical name and wins; `labels`
    are the human headings accepted as a fallback.
    """

    def __init__(self, name, code=None, labels=(), required=True):
        self.name = name
        self.code = code
        self.labels = tuple(labels)
        self.required = required

    @property
    def tokens(self):
        out = []
        if self.code:
            out.append(norm(self.code))
        out.extend(norm(i) for i in self.labels)
        return [t for t in out if t]

    def describe(self):
        bits = [self.code] if self.code else []
        bits.extend(self.labels)
        return " / ".join(bits)


class SheetSpec:
    def __init__(self, key, aliases, fields):
        self.key = key
        self.aliases = tuple(norm(a) for a in aliases)
        self.fields = fields

    def matches(self, sheetname):
        return norm(sheetname) in self.aliases


# How far down a sheet we are willing to look for the header block. The item
# sheet of MASTER DATA.XLSX carries a twelve row legend of material type codes
# before its headers, so a shallow window is not enough. Matching is on exact
# normalised tokens, so scanning into the data rows is harmless: a value like
# "ITEM326404" never equals "codeartic".
HEADER_SEARCH_ROWS = 20


SHEETS = [
    SheetSpec(
        "sections",
        ["9.Sections within plant", "Sections within plant", "sections"],
        [
            Field("code", "CODESECTI", ["Section ID"]),
            Field("name", "DESIGSECT", ["Section designation"]),
        ],
    ),
    SheetSpec(
        "groups",
        ["5.Machine groupings", "Machine groupings", "machinegroups"],
        [
            Field("code", "ILOT", ["Machine grouping"]),
            Field("name", "LIBILOT", ["Machine grouping designation"]),
            Field("mode", "MODEILOT", ["Operating mode"], required=False),
        ],
    ),
    SheetSpec(
        "machines",
        ["6.Machines list", "Machines list", "machines"],
        [
            Field("group", "ILOT", ["Machine grouping"]),
            Field("machine", "MACHINE", ["Machine"]),
            Field("section", "CODESECTI", ["Section ID"], required=False),
            Field("calendar", "NOCALHEBD", ["Calendar ID"], required=False),
            Field("name", "LIBMACH", ["Machine designation"], required=False),
        ],
    ),
    SheetSpec(
        "operations",
        ["3.list of Operations", "list of Operations", "operations"],
        [
            Field("op", "OPE", ["Operation"]),
            Field("group", "ILOT", ["Machine grouping"]),
            Field("qty", "CODEBASET", ["Operation Qty"], required=False),
            Field("cycle", "DURREAL", ["Operation Cycle time"], required=False),
            Field("name", "LIBOP", ["Operation designation"], required=False),
        ],
    ),
    SheetSpec(
        "phases",
        ["7.Phases of the routings", "Phases of the routings", "phases"],
        [
            Field("routing", "NOMG", ["Routing ID"]),
            # Both of the next two are labelled "Phase code" in the human
            # header row, so only the LN code can tell them apart.
            Field("seq", "NOPHASE"),
            Field("op", "OPE"),
            Field("desc", "LIBPHASE", ["Phase designation"], required=False),
            Field("cycle", None, ["Cycle Time"], required=False),
            Field("setup", None, ["Setup Time"], required=False),
            Field("next", None, ["Next Operation"], required=False),
            Field("subcontract", None, ["Subcontracting"], required=False),
        ],
    ),
    SheetSpec(
        "items",
        ["2.List of all the items", "List of all the items", "items"],
        [
            Field("code", "CODEARTIC", ["Item"]),
            Field("uom", "UTIL", ["Unit of measurement"], required=False),
            Field("desc", "LIBARTIC", ["Designation/Description"], required=False),
            Field("type", "TYPEMATI", ["Material type"], required=False),
        ],
    ),
    SheetSpec(
        "workorders",
        ["11.Work order list", "Work order list", "workorders"],
        [
            Field("wo", "NOF", ["Work order id"]),
            Field("routing", "NOMG", ["Routing ID"]),
            Field("item", "CODEARTIC", ["Item"]),
            Field("uom", "UTIL", ["Unit of measurement"], required=False),
            # Present in APS Sample Data.xlsx, absent from MASTER DATA.XLSX.
            # When present it is real data and must beat any generated value.
            Field("qty", "QTY", required=False),
            Field("due", "FPLUSTARD", ["Customer due date"], required=False),
            Field("start", "DPLUSTOT", ["start date"], required=False),
        ],
    ),
]


class SheetReader:
    """
    One resolved sheet: which worksheet it is, where its header block ended,
    and which column index each field lives in.
    """

    def __init__(self, spec, sheetname, columns, first_data_row):
        self.spec = spec
        self.sheetname = sheetname
        self.columns = columns  # field name -> column index
        self.first_data_row = first_data_row
        self.rows = []

    @property
    def key(self):
        return self.spec.key

    def missing(self):
        return [
            f.name for f in self.spec.fields if f.required and f.name not in self.columns
        ]

    def found(self):
        return sorted(self.columns)

    def get(self, row, field, default=None):
        idx = self.columns.get(field)
        if idx is None or idx >= len(row):
            return default
        value = row[idx]
        if value is None:
            return default
        if isinstance(value, str):
            value = value.strip()
            if value == "":
                return default
        return value

    def text(self, row, field, default=""):
        value = self.get(row, field)
        return default if value is None else str(value).strip()


def locateHeader(worksheet_rows, spec):
    """
    Find each field's column by scanning the first rows for its LN code, and
    only then for its human label. Returns (columns, first_data_row).

    The LN code pass runs over every candidate row before the label pass, so a
    duplicated human label can never shadow a correct technical code.
    """
    columns = {}
    header_rows = set()

    for use_code_only in (True, False):
        for row_idx, row in enumerate(worksheet_rows[:HEADER_SEARCH_ROWS]):
            for col_idx, cell in enumerate(row):
                token = norm(cell)
                if not token:
                    continue
                for field in spec.fields:
                    if field.name in columns:
                        continue
                    if use_code_only:
                        if not field.code or token != norm(field.code):
                            continue
                    else:
                        if token not in [norm(i) for i in field.labels]:
                            continue
                    columns[field.name] = col_idx
                    header_rows.add(row_idx)

    first_data_row = (max(header_rows) + 1) if header_rows else 0
    return columns, first_data_row


class WorkbookProblem(Exception):
    pass


class SiegerWorkbook:
    """
    Resolve a planning workbook into frePPLe-shaped rows.

    Usage:
        wbk = SiegerWorkbook(path_or_filelike)
        wbk.read()
        report = wbk.report()      # what was found, what is missing
        if not report["blocking"]:
            payload = wbk.build()
    """

    def __init__(self, source):
        self.source = source
        self.sheets = {}  # key -> SheetReader
        self.unmatched = []  # sheet names we ignored
        self.generated = []

    # -- reading ---------------------------------------------------------

    def read(self):
        from openpyxl import load_workbook

        try:
            wb = load_workbook(filename=self.source, read_only=True, data_only=True)
        except Exception as e:
            raise WorkbookProblem("The file could not be read as .xlsx: %s" % e)

        try:
            for sheetname in wb.sheetnames:
                spec = next((s for s in SHEETS if s.matches(sheetname)), None)
                if not spec:
                    self.unmatched.append(sheetname)
                    continue
                rows = [r for r in wb[sheetname].iter_rows(values_only=True)]
                columns, first_data_row = locateHeader(rows, spec)
                reader = SheetReader(spec, sheetname, columns, first_data_row)
                key_field = spec.fields[0].name
                for row in rows[first_data_row:]:
                    if all(c is None or str(c).strip() == "" for c in row):
                        continue
                    if reader.get(row, key_field) is None:
                        continue
                    reader.rows.append(row)
                self.sheets[spec.key] = reader
        finally:
            wb.close()
        return self

    # -- reporting -------------------------------------------------------

    def report(self):
        """
        A planner-readable account of what the workbook contains. `blocking`
        lists the problems that stop the load; `notes` are informational.
        """
        blocking, notes = [], []
        sheet_rows = []

        for spec in SHEETS:
            reader = self.sheets.get(spec.key)
            if not reader:
                blocking.append(
                    "Sheet for '%s' not found. Expected a tab named like: %s"
                    % (spec.key, ", ".join(spec.aliases))
                )
                sheet_rows.append(
                    {
                        "key": spec.key,
                        "sheet": None,
                        "rows": 0,
                        "header_row": None,
                        "found": [],
                        "missing": [f.name for f in spec.fields if f.required],
                    }
                )
                continue
            missing = reader.missing()
            if missing:
                for name in missing:
                    field = next(f for f in spec.fields if f.name == name)
                    blocking.append(
                        "Sheet '%s': could not find the column for '%s' (looked for %s)"
                        % (reader.sheetname, name, field.describe())
                    )
            if not reader.rows:
                notes.append("Sheet '%s' has no data rows." % reader.sheetname)
            sheet_rows.append(
                {
                    "key": spec.key,
                    "sheet": reader.sheetname,
                    "rows": len(reader.rows),
                    "header_row": reader.first_data_row,
                    "found": reader.found(),
                    "missing": missing,
                }
            )

        wo = self.sheets.get("workorders")
        if wo and "qty" in wo.columns:
            notes.append(
                "Work order quantities will be read from the workbook "
                "(column QTY found) instead of being generated."
            )
        elif wo:
            notes.append(
                "No QTY column in the work order sheet - quantities will be "
                "generated and listed in the generated-value register."
            )

        if self.unmatched:
            notes.append("Ignored sheets: %s" % ", ".join(self.unmatched))

        return {"blocking": blocking, "notes": notes, "sheets": sheet_rows}

    # -- transformation --------------------------------------------------

    def synth(self, category, what, value, why):
        self.generated.append(
            {"category": category, "what": what, "value": value, "why": why}
        )

    def routingPhases(self):
        """Phases grouped by routing code, ordered by sequence."""
        phases = self.sheets["phases"]
        out = defaultdict(list)
        for row in phases.rows:
            out[phases.text(row, "routing")].append(row)
        for key in out:
            out[key].sort(key=lambda r: toInt(phases.get(r, "seq"), 0))
        return out

    def operationGroups(self):
        """Operation code -> the machine groups that can perform it."""
        ops = self.sheets["operations"]
        out = defaultdict(list)
        for row in ops.rows:
            group = ops.text(row, "group")
            if group:
                out[ops.text(row, "op")].append(group)
        return out


def toInt(value, default=0):
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return default


def toFloat(value, default=0.0):
    try:
        return float(value)
    except (TypeError, ValueError):
        return default
