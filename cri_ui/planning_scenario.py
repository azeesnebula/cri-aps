"""Deterministic pump-factory planning example. No real customer data.

Product families are supported by C.R.I's public catalogue; every SKU, BOM,
capacity, cycle time, purchase commitment and customer below is invented.
All datetimes are local Asia/Kolkata, as used by this frePPLe installation.
"""
from collections import defaultdict
from datetime import datetime, time, timedelta
from decimal import Decimal
import json
import math
from pathlib import Path

from django.conf import settings
from django.core.management.base import CommandError
from django.db import transaction

from freppledb.common.models import Parameter
from freppledb.input.models import (
    Buffer, Calendar, CalendarBucket, Customer, Demand, Item, ItemSupplier,
    Location, Operation, OperationMaterial, OperationPlan, OperationResource,
    Resource, Supplier,
)

SOURCE = "CRI-PLANNING-20261019-v1"
START = datetime(2026, 10, 19, 6)
PLANT = "CRI-SYN-COIMBATORE"
SHIFTS = ((6, 0, 10, 0), (10, 30, 14, 0), (14, 0, 18, 0), (18, 30, 22, 0))
SOURCES = [
    ("Product range", "https://www.crigroups.com/pumps/"),
    ("Mini pumps and published materials", "https://www.crigroups.com/pumps/mini-pumps/"),
    ("Borewell construction and motor", "https://www.crigroups.com/borewell-submersible-pumps/"),
    ("Openwell pumps", "https://www.crigroups.com/pump/openwell-submersible-pumps/"),
    ("Monoblock pumps", "https://www.crigroups.com/pump/centrifugal-monoblock-pumps/"),
    ("Vertical multistage pumps", "https://www.crigroups.com/pump/vertical-multistage-inline-pump/"),
]

# code, descriptive planning variant (NOT a commercial model), power kW,
# multistage count, cycle-time multiplier, winding-wire kg, shaft-stock kg
PRODUCTS = (
    ("MINI", "Mini self-priming pump | 0.75 kW", .75, 1, .8, .9, .6),
    ("BORE", "100 mm borewell submersible | 2.2 kW, 12 stages", 2.2, 12, 1.1, 1.8, 1.4),
    ("OPEN", "Openwell submersible pump | 3.7 kW", 3.7, 1, 1.2, 2.4, 1.7),
    ("MONO", "Centrifugal monoblock pump | 3.7 kW", 3.7, 1, 1.0, 2.2, 1.5),
    ("VM", "Vertical multistage inline | 5.5 kW, 6 stages", 5.5, 6, 1.5, 3.1, 2.1),
)

# Due is the ex-factory dispatch cutoff, NOT arrival at the customer's city.
# Two lines belong to each all-together shipment; no internal component demand.
SHIPMENTS = (
    ("Coimbatore residential distributor", "2026-10-28", (("MINI", 60), ("MONO", 20))),
    ("Erode agricultural distributor", "2026-10-30", (("BORE", 40), ("OPEN", 30))),
    ("Chennai housing contractor", "2026-11-03", (("MINI", 70), ("MONO", 20))),
    ("Madurai irrigation project", "2026-11-05", (("BORE", 40), ("OPEN", 30))),
    ("Bengaluru water treatment integrator", "2026-11-07", (("VM", 20), ("MINI", 20))),
    ("Salem agricultural distributor", "2026-11-14", (("OPEN", 40), ("MONO", 30))),
    ("Kochi commercial buildings", "2026-11-17", (("VM", 20), ("BORE", 30))),
    ("Hyderabad residential distributor", "2026-11-19", (("MINI", 60), ("MONO", 30))),
    ("Pune process-water integrator", "2026-11-21", (("VM", 20), ("OPEN", 30))),
    ("Tiruppur industrial water project", "2026-11-24", (("BORE", 40), ("VM", 20))),
)

# Group, section, staffed machines, assumed standard hourly cost in INR.
WORKCENTRES = (
    ("TURN", "MACHINING", 1, 650), ("GRIND", "MACHINING", 1, 550),
    ("CNC", "MACHINING", 1, 850), ("HYD", "MACHINING", 1, 600),
    ("WIND", "MOTORS", 1, 450), ("ELEC", "MOTORS", 1, 400),
    ("MOTOR", "MOTORS", 1, 450), ("ASSY", "ASSEMBLY", 2, 400),
    ("TEST", "QUALITY", 2, 500), ("PACK", "DISPATCH", 1, 300),
)
MAINTENANCE = {
    "CRI-TEST-01": (datetime(2026, 10, 28, 14), datetime(2026, 10, 28, 18)),
    "CRI-CNC-01": (datetime(2026, 11, 3, 6), datetime(2026, 11, 3, 10)),
}
INPUT_MODELS = (
    Calendar, CalendarBucket, Location, Customer, Supplier, Resource, Item,
    Operation, OperationMaterial, OperationResource, ItemSupplier, Buffer, Demand,
)


def scalar(value):
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, timedelta):
        return value.total_seconds()
    if isinstance(value, (datetime, time)):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    return value


def json_default(value):
    converted = scalar(value)
    if converted is value:
        raise TypeError(type(value).__name__)
    return converted


def check_database(db):
    if settings.DATABASES[db]["NAME"] != "cri0":
        raise CommandError("This planning scenario may only be loaded/exported on the C.R.I cri0 database.")


def load(db="default"):
    check_database(db)
    # Never empty, replace, or mingle an existing company's planning inputs.
    if any(m.objects.using(db).exists() for m in (Item, Demand, Resource, Operation, OperationPlan)):
        raise CommandError("Planning data already exists. This loader only accepts an empty C.R.I dataset; use export or replan instead.")
    with transaction.atomic(using=db):
        def create(model, **fields):
            return model.objects.using(db).create(source=SOURCE, **fields)

        def parameter(name, value, description):
            Parameter.objects.using(db).update_or_create(name=name, defaults={
                "value": value, "description": description, "source": SOURCE})

        parameter("currentdate", START.isoformat(sep=" "), "Plan start in Asia/Kolkata")
        parameter("cri.planning.dataset", SOURCE, "Illustrative planning inputs; see scenario workbook")
        parameter("plan.loglevel", "1", "Planning scenario: normal solver diagnostics")
        parameter("plan.deliveryDuration", "0", "Dispatch due dates are ex-factory; no transit time is claimed")
        parameter("currency", "INR", "Planning costs in INR")
        parameter("sieger.workdays", "0,1,2,3,4,5", "Monday-Saturday working week")

        def calendar(name, maintenance=None):
            cal = create(Calendar, name=name, defaultvalue=0,
                         description="Mon-Sat 06-14 / 14-22; 10:00 and 18:00 half-hour breaks")
            for index, (h1, m1, h2, m2) in enumerate(SHIFTS):
                # CalendarBucket endtime is inclusive to the second in frePPLe.
                exclusive_end = datetime(2026, 1, 1, h2, m2) - timedelta(seconds=1)
                create(CalendarBucket, calendar=cal, startdate=START,
                       enddate=datetime(2027, 1, 1), priority=10 + index,
                       value=1, sunday=False, starttime=time(h1, m1), endtime=exclusive_end.time())
            if maintenance:
                create(CalendarBucket, calendar=cal, startdate=maintenance[0],
                       enddate=maintenance[1], priority=-1, value=0)
            return cal

        shifts = calendar("CRI-SYN-2SHIFT")
        location = create(Location, name=PLANT,
                          description="Coimbatore pump plant; illustrative facility layout")
        sections = {name: create(Resource, name=f"CRI-SEC-{name}", description=name.title(),
                                type="infinite", location=location)
                    for name in sorted({w[1] for w in WORKCENTRES})}
        groups = {}
        for group, section, count, cost in WORKCENTRES:
            groups[group] = create(Resource, name=f"CRI-GRP-{group}", description=f"{group} staffed work centre",
                                   owner=sections[section], type="default", location=location,
                                   maximum=1, constrained=True, available=shifts, efficiency=85)
            for index in range(1, count + 1):
                name = f"CRI-{group}-{index:02d}"
                cal = calendar(name + "-CAL", MAINTENANCE[name]) if name in MAINTENANCE else shifts
                create(Resource, name=name, description=f"Staffed {group.lower()} cell {index}",
                       owner=groups[group], location=location, type="default", maximum=1,
                       constrained=True, available=cal, efficiency=85, cost=cost,
                       maxearly=timedelta(days=60))

        suppliers = {}
        for code, desc in (("METALS", "Castings and stainless bar"), ("ELECTRICAL", "Copper wire and laminations"),
                           ("HYDRAULICS", "Impellers and hydraulic stage components"),
                           ("STANDARD", "Bearings, seals and hardware"), ("PACKAGING", "Cartons and crates")):
            suppliers[code] = create(Supplier, name=f"SYN-SUP-{code}", description=f"Fictitious supplier: {desc}")

        items, raw, boms = {}, {}, {}

        def item(name, description, category, cost, uom="each"):
            obj = create(Item, name=name, description=description, category=category, uom=uom,
                         type="make to stock", cost=cost, subcategory="")
            items[name] = obj
            return obj

        def buy(name, description, supplier, days, cost, uom="each", multiple=1):
            obj = item(name, description, "PURCHASED", cost, uom)
            create(ItemSupplier, item=obj, location=location, supplier=suppliers[supplier],
                   leadtime=timedelta(days=days), sizeminimum=multiple, sizemultiple=multiple,
                   batchwindow=timedelta(0), cost=cost, priority=1)
            raw[name] = {"supplier": suppliers[supplier], "days": days, "multiple": multiple}
            return obj

        buy("SYN-SS410-BAR", "SS410 shaft bar; kg", "METALS", 4, 180, "kg", 5)
        buy("SYN-COPPER-ENAMEL", "Enamelled copper winding wire; kg", "ELECTRICAL", 5, 880, "kg", 5)
        buy("SYN-COPPER-WATERPROOF", "Waterproof insulated submersible winding wire; kg", "ELECTRICAL", 6, 960, "kg", 5)
        buy("SYN-HARDWARE", "Fastener and O-ring kit", "STANDARD", 3, 90)
        buy("SYN-LEAD-CABLE", "Electrical lead cable; metres", "ELECTRICAL", 3, 70, "m", 10)
        buy("SYN-THERMAL-PROTECTOR", "Thermal overload protector", "ELECTRICAL", 4, 120)

        def make(name, description, category, inputs, steps, multiplier, lot=100):
            """A routing consumes its BOM at its first step and produces at its end.

            Subassemblies carry their own independent routings; only FG has demand.
            Fixed duration is a conservative setup per lot, not a sequence matrix.
            """
            obj = item(name, description, category, 1000)
            parent = create(Operation, name=f"RT-{name}", description=description,
                            item=obj, location=location, type="routing", priority=1,
                            sizeminimum=1, sizemultiple=1, sizemaximum=lot,
                            category=category, subcategory="", batchwindow=timedelta(0))
            for index, (group, label, fixed_min, per_min) in enumerate(steps, 1):
                step = create(Operation, name=f"{name}-{index * 10:03d}-{group}",
                              description=label, owner=parent, priority=index * 10, type="time_per",
                              location=location, duration=timedelta(minutes=fixed_min),
                              duration_per=timedelta(minutes=round(per_min * multiplier, 2)),
                              available=shifts, sizeminimum=1, sizemultiple=1, category=category,
                              subcategory="")
                create(OperationResource, operation=step, resource=groups[group], quantity=1,
                       search="MINCOST", priority=1)
                if index == 1:
                    for component, qty in inputs.items():
                        create(OperationMaterial, operation=step, item=items[component],
                               quantity=-Decimal(str(qty)), type="start")
            # Output belongs to the last routing step, not the routing wrapper.
            create(OperationMaterial, operation=step, item=obj, quantity=1, type="end")
            boms[name] = inputs
            return obj

        for code, desc, kw, stages, factor, copper, steel in PRODUCTS:
            prefix = f"SYN-{code}"
            housing = "Forged brass casing blank" if code == "MINI" else "Stainless casing blank" if code in ("BORE", "VM") else "Cast-iron casing blank"
            buy(prefix + "-CASE-RAW", housing, "METALS", 5, 900 * factor)
            buy(prefix + "-IMPELLER-RAW", "High-tensile brass impeller blank" if code == "MINI" else "Hydraulic impeller blank", "HYDRAULICS", 6, 170 * factor)
            buy(prefix + "-DIFFUSER", "Diffuser / guide vane set", "HYDRAULICS", 6, 110 * factor)
            buy(prefix + "-LAMINATION", "Motor stator lamination stack", "ELECTRICAL", 5, 450 * factor)
            buy(prefix + "-ROTOR", "Rotor blank and motor shell kit", "ELECTRICAL", 5, 550 * factor)
            buy(prefix + "-BEARING", "Matched motor bearing", "STANDARD", 3, 130 * factor)
            buy(prefix + "-SEAL", "Compatible mechanical seal kit", "STANDARD", 4, 180 * factor)
            buy(prefix + "-PACK", "Protective carton/crate and labels", "PACKAGING", 2, 150 * factor)

            make(prefix + "-SHAFT", "Machined and ground shaft: " + code, "SHAFT",
                 {"SYN-SS410-BAR": round(steel * 1.03, 4)},
                 [("TURN", "Shaft turning and threading", 25, 6), ("GRIND", "Shaft grinding and dimensional inspection", 15, 4)], factor)
            make(prefix + "-CASING", "Machined casing: " + code, "CASING",
                 {prefix + "-CASE-RAW": 1},
                 [("CNC", "Casing bore, flange and port machining", 30, 12), ("HYD", "Deburr, clean and inspect casing", 10, 4)], factor)
            make(prefix + "-HYDRAULIC", "Balanced hydraulic cartridge: " + code, "HYDRAULIC",
                 {prefix + "-IMPELLER-RAW": stages, prefix + "-DIFFUSER": stages},
                 [("HYD", "Finish and balance hydraulic components", 20, 7 + stages * .6),
                  ("HYD", "Build and inspect hydraulic cartridge", 10, 5 + stages * .5)], factor)
            wire = "SYN-COPPER-WATERPROOF" if code in ("BORE", "OPEN") else "SYN-COPPER-ENAMEL"
            make(prefix + "-STATOR", "Wound and tested stator: " + code, "STATOR",
                 {prefix + "-LAMINATION": 1, wire: round(copper * 1.02, 4)},
                 [("WIND", "Wind, insert and terminate stator coils", 30, 20),
                  ("ELEC", "Insulation, surge and winding resistance test", 10, 8)], factor)
            make(prefix + "-MOTOR", "Assembled and electrically tested motor: " + code, "MOTOR",
                 {prefix + "-STATOR": 1, prefix + "-SHAFT": 1, prefix + "-ROTOR": 1,
                  prefix + "-BEARING": 2, "SYN-LEAD-CABLE": 3, "SYN-THERMAL-PROTECTOR": 1},
                 [("MOTOR", "Motor assembly, alignment and rotor fit", 20, 12),
                  ("ELEC", "Motor electrical safety and no-load test", 10, 6)], factor)
            make(prefix + "-PUMP", desc, "FINISHED PUMP",
                 {prefix + "-MOTOR": 1, prefix + "-CASING": 1, prefix + "-HYDRAULIC": 1,
                  prefix + "-SEAL": 1, prefix + "-PACK": 1, "SYN-HARDWARE": 1},
                 [("ASSY", "Final pumpset assembly and leak check", 20, 14),
                  ("TEST", "Hydraulic performance and pressure test", 15, 12),
                  ("PACK", "Final inspection, nameplate, preservation and packing", 10, 6)], factor, lot=100)

        names, demand_totals = {}, defaultdict(int)
        for index, (description, date, lines) in enumerate(SHIPMENTS, 1):
            tag = f"CRI-SHIP-{index:03d}"
            customer = create(Customer, name=f"SYN-CUSTOMER-{index:02d}",
                              description="Fictitious customer: " + description)
            names[tag] = {"desc": description}
            for line, (code, qty) in enumerate(lines, 1):
                name = f"SYN-{code}-PUMP"
                create(Demand, name=f"CRI-SO-{index:03d}-{line:02d}", item=items[name],
                       location=location, customer=customer, quantity=qty, status="open",
                       category=tag, subcategory=f"SYN-PROJECT-{index:02d}", owner=tag,
                       description=description, priority=10,
                       due=datetime.fromisoformat(date + " 17:00:00"),
                       minshipment=qty, policy="alltogether", maxlateness=timedelta(days=30))
                demand_totals[name] += qty
        parameter("sieger.shipnames", json.dumps(names), "Shipment names for legacy-compatible UI")

        # Explode ONLY the independent FG demand. Never duplicate demand at BOM levels.
        gross = defaultdict(Decimal)

        def explode(name, qty):
            if name in raw:
                gross[name] += qty
            else:
                for child, factor in boms[name].items():
                    explode(child, qty * Decimal(str(factor)))

        for name, qty in demand_totals.items():
            explode(name, Decimal(qty))
        for index, (name, info) in enumerate(sorted(raw.items()), 1):
            required = gross[name]
            multiple = Decimal(info["multiple"])
            opening = math.ceil(required * Decimal("0.35") / multiple) * multiple
            receipt = math.ceil(required * Decimal("0.35") / multiple) * multiple
            create(Buffer, item=items[name], location=location, type="default",
                   onhand=opening, minimum=0, description="Counted opening inventory")
            # Previously ordered replenishment. It is confirmed, so planning must respect it.
            arrival = datetime(2026, 10, 22 if info["days"] <= 4 else 23, 8)
            create(OperationPlan, reference=f"CRI-SYN-PO-{index:03d}", type="PO", status="confirmed",
                   item=items[name], supplier=info["supplier"], location=location, quantity=receipt,
                   startdate=arrival - timedelta(days=info["days"]), enddate=arrival,
                   remark="Existing supplier commitment; remaining demand requires proposed POs")
        Resource.rebuildHierarchy(db)
        Item.rebuildHierarchy(db)
        return {"shipments": len(SHIPMENTS), "order_lines": sum(len(s[2]) for s in SHIPMENTS),
                "pumps": sum(demand_totals.values()), "gross_material_requirements": dict(gross)}


def write_workbook(path, sheets):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name[:31])
        if not rows:
            ws.append(["No records"])
            continue
        columns = list(rows[0])
        ws.append(columns)
        for row in rows:
            values = []
            for key in columns:
                value = scalar(row.get(key))
                if isinstance(value, (dict, list, tuple)):
                    value = json.dumps(value, default=json_default)
                values.append(value)
            ws.append(values)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="C8102E")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(wrap_text=True)
        for index, name in enumerate(columns, 1):
            width = max(len(str(name)), *(len(str(row.get(name) or "")) for row in rows[:100]))
            ws.column_dimensions[get_column_letter(index)].width = min(58, max(15, width + 2))
    wb.save(path)


def export_inputs(path, db="default"):
    path = Path(path)
    path.mkdir(parents=True, exist_ok=True)
    tables = {m._meta.db_table: list(m.objects.using(db).filter(source=SOURCE).values()) for m in INPUT_MODELS}
    tables["confirmed_purchase_orders"] = list(OperationPlan.objects.using(db).filter(
        source=SOURCE, type="PO", status="confirmed").values())
    tables["parameters"] = list(Parameter.objects.using(db).filter(source=SOURCE).values())
    metadata = {"dataset": SOURCE, "illustrative_inputs": True, "plan_start": START, "timezone": "Asia/Kolkata",
                "duration_unit": "seconds", "due_basis": "ex-factory dispatch, 17:00 local; no transit included",
                "sources": [{"subject": label, "url": url} for label, url in SOURCES]}
    (path / "planning_inputs.json").write_text(json.dumps({"metadata": metadata, "tables": tables},
                                               default=json_default, indent=2), encoding="utf-8")
    guide = [{"topic": "Dataset", "details": SOURCE},
             {"topic": "Scope", "details": "Illustrative planning inputs, not actual company records. Product families only are based on the public catalogue."},
             {"topic": "Plan start", "details": "19 October 2026 06:00 Asia/Kolkata"},
             {"topic": "Due dates", "details": "Ex-factory dispatch cutoff 17:00, 28 Oct-24 Nov; based on trial capacity results plus two working days and at least two working days between commitments. Customer transit is excluded."},
             {"topic": "Calendar", "details": "Mon-Sat 06-14 / 14-22; breaks 10-10:30 and 18-18:30; Sunday off; no other holidays assumed."},
             {"topic": "Capacity", "details": "12 staffed single-job cells, 85% efficiency. One operator is assumed dedicated to each operating cell per shift."},
             {"topic": "BOM", "details": "One FG demand only; six manufactured levels/items per family; bought-in foundry blanks, rotor kits and stage parts."},
             {"topic": "Yield", "details": "100% piece yield and no rework assumed; 3% steel and 2% wire allowances included in input quantities."},
             {"topic": "Setups", "details": "Fixed setup per lot; all manufactured lots up to 100 to support full-quantity all-together shipments; no sequence-dependent changeovers."},
             {"topic": "Inventory", "details": "35% opening stock, 35% confirmed receipts on 22/23 Oct; balance planned with 2-6 calendar-day supplier lead times."},
             {"topic": "Units", "details": "All durations are seconds; datetime values local Asia/Kolkata; quantities use each item's UOM; monetary assumptions INR."},
             {"topic": "Workbook", "details": "Complete reference snapshot; regenerate with frepplectl criplanning --load. Not the numbered legacy Infor-LN import format."}]
    write_workbook(path / "CRI_Planning_Inputs_2026-10-19.xlsx", {"Read me": guide, "Sources": metadata["sources"], **tables})
    return tables
